import nest_asyncio
from telegram import Update, InputFile, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    ApplicationBuilder,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    CallbackQueryHandler,
    filters,
)
from openpyxl import Workbook, load_workbook

# ================= CONFIG =================
BOT_TOKEN = '8481478480:AAGwfS9t9jb-9qoVsjdRq0mWvvdH7PNoOQM'
WALLET_ADDRESS = '6YyfK7W8GHSYdrKCtYQ9GbBEntSsmHHb9FWzcSvpijYk'
PRODUCT_FILE = 'OnlyFans Mod.apk'
QR_FILE = 'qr.png'  # Your pre-uploaded QR code file
EXPECTED_AMOUNT = 3.0
EXCEL_FILE = 'user_data.xlsx'
ADMIN_ID = 5928995459
# ==========================================

# Setup Excel workbook and sheet
try:
    wb = load_workbook(EXCEL_FILE)
    sheet = wb.active
except FileNotFoundError:
    wb = Workbook()
    sheet = wb.active
    sheet.append(['UserID', 'Username', 'PaymentAmount', 'ScreenshotFileID', 'Status'])
    wb.save(EXCEL_FILE)


def save_new_user(user_id, username, payment_amount, screenshot_file_id=None, status='pending'):
    sheet.append([user_id, username, payment_amount, screenshot_file_id, status])
    wb.save(EXCEL_FILE)


def update_payment_status(user_id, new_status):
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == user_id:
            row[4].value = new_status
            wb.save(EXCEL_FILE)
            break


def update_screenshot_file_id(user_id, file_id):
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == user_id:
            row[3].value = file_id
            wb.save(EXCEL_FILE)
            break


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("👋 Welcome! Use /buy to get payment info.")


async def buy(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user

    msg = (
    f"💸 Send **${EXPECTED_AMOUNT} USDT (SPL)** to this wallet\n\n"
    f"📌 Public Address (hold to copy):\n<code>{WALLET_ADDRESS}</code>\n\n"
    f"📷 Scan the QR code below or copy the address above.\n\n"
    f"⚠️ Important Instructions:\n"
    f"1️⃣ **Network:** Only use **Solana (SPL/USDT)** network. "
    f"Do NOT send regular SOL or USDT on TRC20/ERC20. Sending on wrong network will result in permanent loss of funds.\n"
    f"2️⃣ **Supported Wallets:**\n"
    f"   - Solflare Wallet\n"
    f"   - Trust Wallet (ensure USDT SPL network is selected)\n"
    f"   - Binance Wallet (select USDT SPL/USDT Solana network)\n"
    f"3️⃣ After payment, type /confirm and send a screenshot of your transaction.\n\n"
    f"💡 Tip: In most wallets, you can choose 'USDT SPL' or 'USDT Solana' before sending. Make sure it matches exactly!"
)


    save_new_user(user.id, user.username or 'No username', EXPECTED_AMOUNT, status='pending')

    # Send the pre-uploaded QR image
    with open(QR_FILE, "rb") as qr_image:
        await update.message.reply_photo(
            photo=qr_image,
            caption=msg,
            parse_mode="HTML"
        )


async def confirm(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    update_payment_status(user.id, 'pending')
    await update.message.reply_text(
        "📸 Please upload a screenshot of your transaction.\n\n"
        "⏳ We will verify your payment manually and send the APK when ready."
    )


async def handle_screenshot(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user

    if not update.message.photo:
        await update.message.reply_text("Please send a valid photo.")
        return

    file_id = update.message.photo[-1].file_id
    update_screenshot_file_id(user.id, file_id)

    username_text = f"@{user.username}" if user.username else "No username"

    caption_text = (
        f"📸 Screenshot received\n\n"
        f"👤 Username: {username_text}\n"
        f"🆔 UserID: <code>{user.id}</code>"
    )

    keyboard = InlineKeyboardMarkup([
        [
            InlineKeyboardButton("✅ Mark as Paid", callback_data=f"markpaid:{user.id}"),
            InlineKeyboardButton("❌ Reject Payment", callback_data=f"reject:{user.id}")
        ]
    ])

    await context.bot.send_photo(
        chat_id=ADMIN_ID,
        photo=file_id,
        caption=caption_text,
        parse_mode="HTML",
        reply_markup=keyboard
    )

    # Notify user based on current status
    status = None
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == user.id:
            status = row[4].value
            break

    if status == 'paid':
        await update.message.reply_text("✅ Payment confirmed! You already have the APK.")
    elif status == 'pending':
        await update.message.reply_text("📥 Screenshot received! We will verify your payment manually.")
    else:
        await update.message.reply_text("ℹ️ Please use /buy to get payment info first.")


async def button_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()

    if ":" not in query.data:
        return

    action, user_id_str = query.data.split(":")
    user_id = int(user_id_str)

    if str(query.from_user.id) != str(ADMIN_ID):
        await query.edit_message_caption(
            caption="⛔ You are not authorized to perform this action."
        )
        return

    if action == "markpaid":
        update_payment_status(user_id, 'paid')

        await query.edit_message_caption(
            caption=f"✅ User <code>{user_id}</code> marked as paid.",
            parse_mode="HTML"
        )

        try:
            with open(PRODUCT_FILE, "rb") as f:
                await context.bot.send_document(
                    chat_id=user_id,
                    document=InputFile(f, filename=PRODUCT_FILE),
                    caption="🎉 Your payment has been confirmed!\n📲 Here is your APK file.",
                    write_timeout=120,
                    connect_timeout=120,
                    read_timeout=120,
                    pool_timeout=120
                )

            # Notify admin
            await context.bot.send_message(
                chat_id=ADMIN_ID,
                text=f"App delivered to User <code>{user_id}</code>",
                parse_mode="HTML"
            )

        except Exception as e:
            await context.bot.send_message(
                chat_id=ADMIN_ID,
                text=f"⚠️ Failed to send APK to user {user_id}: {e}"
            )

    elif action == "reject":
        update_payment_status(user_id, 'rejected')

        await query.edit_message_caption(
            caption=f"❌ User <code>{user_id}</code>'s payment has been rejected.",
            parse_mode="HTML"
        )

        await context.bot.send_message(
            chat_id=user_id,
            text="⚠️ Your payment was not verified. Please try again or contact support."
        )


async def main():
    app = ApplicationBuilder().token(BOT_TOKEN).build()

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("buy", buy))
    app.add_handler(CommandHandler("confirm", confirm))
    app.add_handler(MessageHandler(filters.PHOTO & (~filters.COMMAND), handle_screenshot))
    app.add_handler(CallbackQueryHandler(button_handler))

    print("✅ Bot starting...")
    await app.run_polling()


if __name__ == "__main__":
    nest_asyncio.apply()
    import asyncio
    asyncio.get_event_loop().run_until_complete(main())
